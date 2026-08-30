"""Retrofit an existing controlled document into house style.

    uv run --with python-docx --with openpyxl python \
        .claude/skills/gxp-doc/scripts/restyle_doc.py docs_qms/*.docx \
        -o build/qms-restyled --report build/qms-gap-report.md

Reads the source and re-emits its content through gxpdoc, rather than
patching OXML in place: there is then exactly one code path that decides how
a document looks, so a retrofitted file and a newly scaffolded one cannot
drift apart.

NEVER writes to the source. The QMS directories are symlinks into Google
Drive and hold approved controlled documents; output goes to --outdir for
review, and promoting it is a change-control decision, not this script's.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).parent))

import check_doc as lint  # noqa: E402
from gxpdoc import BANNER, COMPANY_PLACEHOLDER, GxpDoc, VERSION_RE  # noqa: E402
from new_doc import slug  # noqa: E402

MASTER_LIST = Path("docs_qms/Document_Master_List.xlsx")

SPECIAL_SECTIONS = {"REVISION HISTORY", "APPROVAL", "REFERENCES"}

#: Section name -> the headers a Label: value run should be given there.
LABEL_TABLE_HEADERS = {
    "RESPONSIBILITIES": ["Role", "Responsibilities"],
    "DEFINITIONS": ["Term", "Definition"],
    "FORMS AND LOGS": ["Document ID", "Purpose"],
    "FORMS AND WORK INSTRUCTIONS": ["Document ID", "Purpose"],
}

#: Prose asserting the old draft-versioning convention, which the house style
#: replaces. Flagged for a human because it is content, not formatting.
STALE_DRAFT_RE = re.compile(r"version[^.]{0,40}\bDRAFT\b", re.IGNORECASE)
HEADING_NUMBER_RE = re.compile(r"^\s*(\d+(?:\.\d+)*)\.?\s+(.*)$", re.DOTALL)


def load_master_list(path=MASTER_LIST):
    """Document ID -> Title, so References rows can be filled in."""
    try:
        import openpyxl
    except ImportError:
        return {}
    if not Path(path).exists():
        return {}
    book = openpyxl.load_workbook(path, read_only=True)
    titles = {}
    for sheet in book:
        for row in sheet.iter_rows(values_only=True):
            if not row or not row[0]:
                continue
            key = str(row[0]).strip()
            if lint.ID_RE.fullmatch(key) and len(row) > 1 and row[1]:
                titles[key] = str(row[1]).strip()
    return titles


def source_sections(ordered, title=""):
    """Walk the source into (section name, [blocks]) in document order.

    A leading heading that merely restates the document title is dropped --
    the repeating page header carries it now.
    """
    sections, current, preamble = [], None, []
    for kind, block in ordered:
        if kind == "p" and lint.is_level1(block):
            name = lint.heading_text(block)
            if not name:
                continue  # SOP-013's stray empty headings
            if not sections and title and name == title.strip().upper():
                continue
            current = (name, [])
            sections.append(current)
            continue
        (current[1] if current else preamble).append((kind, block))
    return preamble, sections


def is_bullet(paragraph):
    name = lint.style_name(paragraph)
    return "List" in name or paragraph.text.strip().startswith(("•", "-", "●"))


def emit_label_table(out, run, section):
    rows = [tuple(p.split(":", 1)) for p in run]
    headers = LABEL_TABLE_HEADERS.get(section, ["Item", "Detail"])
    out.table(headers, [(a.strip(), b.strip()) for a, b in rows])


def emit_blocks(out, blocks, notes, manual, section=""):
    """Re-emit one section's content in house style.

    Prose sitting under a numbered step is part of that step, so it becomes a
    clause one level deeper -- that is what makes it citable as §6.1.1 rather
    than "the third paragraph of 6.1".
    """
    pending, state = [], {"depth": None}

    def emit_prose(text):
        if state["depth"] is None:
            out.body(text)
        else:
            out.clause(min(state["depth"] + 1, 4), text)

    def flush():
        if len(pending) >= 3:
            emit_label_table(out, pending, section)
            notes.append(f"converted {len(pending)} 'Label: value' paragraphs "
                         f"into a table (first: {pending[0][:50]!r})")
        else:
            for text in pending:
                emit_prose(text)
        pending.clear()

    for kind, block in blocks:
        if kind == "t":
            flush()
            grid = lint.rows_of(block)
            if not grid:
                continue
            out.table(grid[0], grid[1:] or [tuple("" for _ in grid[0])])
            continue

        text = block.text.strip()
        if not text:
            continue
        match = HEADING_NUMBER_RE.match(text)
        style = lint.style_name(block)
        if match and style.startswith(("Heading", "GxP L")):
            flush()
            level = max(min(match.group(1).count(".") + 1, 4), 2)
            out.clause(level, match.group(2).strip())
            state["depth"] = level
            continue
        if STALE_DRAFT_RE.search(text):
            manual.append("prose still describes DRAFT as a version value — "
                          f"rewrite: {text[:80]!r}")
        if lint.LABEL_RE.match(text):
            pending.append(text)
            continue
        flush()
        if is_bullet(block):
            out.bullet(text.lstrip("•-● ").strip())
        else:
            emit_prose(text)
    flush()


def collect(sections, name):
    for section_name, blocks in sections:
        if section_name == name:
            return blocks
    return []


def first_table(blocks):
    for kind, block in blocks:
        if kind == "t":
            return lint.rows_of(block)
    return []


def restyle(path, outdir, titles, company=COMPANY_PLACEHOLDER):
    path, notes, manual = Path(path), [], []
    doc = Document(str(path))
    ordered = list(lint.blocks(doc.element.body, doc))

    control = lint.find_control_block(ordered)
    fields = {}
    if control is not None:
        grid = lint.rows_of(control)
        body = grid[1:] if [c.lower() for c in grid[0][:2]] in (
            ["field", "details"], ["field", "value"]) else grid
        fields = {r[0].strip(): r[1].strip() for r in body if r[0].strip()}

    if fields:
        carried = ", ".join(f"{k}: {v!r}" for k, v in fields.items()
                            if k not in ("Document ID", "Title", "Version"))
        notes.append("removed the metadata block; the page header now carries "
                     "ID, description and version")
        if carried:
            manual.append("metadata that was in the removed block, for the "
                          f"Document Master List if not already there — {carried}")

    doc_id = fields.get("Document ID", "") or path.name.split("_")[0]
    title = fields.get("Title", "")
    if not title:
        title = re.sub(r"[_-]v?\d+([._]\d+)*$", "", path.stem).replace("_", " ")
        manual.append(f"no Title row in the control block; derived "
                      f"{title!r} from the filename — set a real title")
    version = fields.get("Version", "") or "1.0"
    if not VERSION_RE.match(version):
        notes.append(f"version was {version!r}; set to 1.0 — a version is never "
                     "the word DRAFT")
        version = "1.0"

    prefix = doc_id.split("-")[0].upper()
    doc_type = lint.PREFIX_TO_TYPE.get(prefix, "sop")
    if doc_type not in BANNER:
        doc_type = "sop"

    preamble, sections = source_sections(ordered, title)
    out = GxpDoc(doc_id=doc_id, title=title, doc_type=doc_type,
                 version=version, company=company or fields.get("Company")
                 or COMPANY_PLACEHOLDER)
    out.toc()

    old_numbers, new_numbers = [], []
    for name, blocks in sections:
        if name in SPECIAL_SECTIONS:
            continue
        old_numbers.append(name)
        out.h1(name.title())
        new_numbers.append(f"{len(new_numbers) + 1}. {name}")
        emit_blocks(out, blocks, notes, manual, name)

    # References -- seeded from what the body actually cites.
    cited = sorted({
        m.group(0)
        for kind, block in ordered if kind == "p"
        for m in lint.ID_RE.finditer(block.text)
        if not lint.PLACEHOLDER_ID_RE.match(m.group(0).split("-", 1)[-1])
    })
    existing = {r[0].strip(): r[1].strip()
                for r in first_table(collect(sections, "REFERENCES"))[1:]
                if len(r) > 1 and r[0].strip()}
    rows, unresolved = [], []
    for identifier in cited:
        name = existing.get(identifier) or titles.get(identifier)
        if not name:
            name = "<title — resolve from the Document Master List>"
            unresolved.append(identifier)
        rows.append((identifier, name))
    if unresolved:
        manual.append("References rows needing a title: " + ", ".join(unresolved))
    out.references({"References": rows})

    history = first_table(collect(sections, "REVISION HISTORY"))
    history_rows = []
    for row in history[1:]:
        padded = list(row) + [""] * (5 - len(row))
        history_rows.append(tuple(padded[:5]))
    if not history_rows:
        history_rows = [(version, "", "Quality Team", "Initial release", "")]
    if len(history) > 1 and len(history[0]) < 5:
        notes.append(f"Revision History had {len(history[0])} columns; padded to "
                     "Version | Date | Author | Description of Change | "
                     "Change Control Ref")
    out.revision_history(history_rows)

    approval = first_table(collect(sections, "APPROVAL"))
    signatories = []
    for row in approval[1:]:
        padded = list(row) + [""] * 3
        signatories.append((padded[0], padded[1], padded[2]))
    if not signatories:
        signatories = [("Reviewer", "<Name>", "<Title>"),
                       ("Approver", "<Name>", "<Title>")]
    if approval and [c.strip() for c in approval[0]][:3] == ["Role", "Name",
                                                             "Signature"]:
        notes.append("Approval table had no Title column; added one")
    out.approval(signatories)

    name = f"{doc_id}_{slug(title)}_v{version}_DRAFT1.docx"
    written = out.save(Path(outdir) / name)
    manual.append("inline bold/italic and any images in the source are not "
                  "carried over; check the output against the original")
    return written, notes, manual


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("paths", nargs="+", type=Path)
    parser.add_argument("-o", "--outdir", type=Path,
                        default=Path("build/qms-restyled"))
    parser.add_argument("--report", type=Path,
                        default=Path("build/qms-gap-report.md"))
    parser.add_argument("--company", default="",
                        help="header wordmark; defaults to the Company row of "
                             "the document's own metadata block, if it has one")
    parser.add_argument("--master-list", type=Path, default=MASTER_LIST,
                        help="xlsx used to resolve References titles")
    args = parser.parse_args()

    titles = load_master_list(args.master_list)
    args.outdir.mkdir(parents=True, exist_ok=True)
    report = ["# QMS retrofit gap report", "",
              "Candidate files are in `%s`. Nothing was written to Google "
              "Drive; promoting these is a change-control decision under "
              "SOP-003 (Change Control)." % args.outdir, ""]

    for path in args.paths:
        if path.name.startswith("~$"):
            continue
        report.append(f"## {path.name}")
        try:
            written, notes, manual = restyle(path, args.outdir, titles,
                                             args.company)
        except Exception as exc:
            report += [f"- **could not restyle**: {exc}", ""]
            print(f"FAILED {path.name}: {exc}")
            continue
        print(f"wrote {written}")
        report.append(f"Output: `{written}`")
        report.append("")
        report.append("### Applied automatically")
        for note in notes:
            report.append(f"- {note}")
        remaining = lint.check(written)
        report += ["", "### Needs a human"]
        for note in manual:
            report.append(f"- {note}")
        for item in remaining.items:
            where = f" — `{item['where']}`" if item["where"] else ""
            report.append(f"- **{item['rule']}** {item['message']}{where}")
        if not manual and not remaining.items:
            report.append("- nothing; the output lints clean")
        report.append("")

    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text("\n".join(report))
    print(f"\nreport: {args.report}")


if __name__ == "__main__":
    main()
